import io
from datetime import datetime
from docx import Document
from fpdf import FPDF

class ExportService:
    @staticmethod
    def generate_docx(synthesis_text: str, schema_title: str = "Occo AI Compliance Audit") -> io.BytesIO:
        doc = Document()
        doc.add_heading(schema_title, level=1)
        doc.add_paragraph(f"Generated Timestamp: {datetime.utcnow().isoformat()} UTC")
        doc.add_heading("LLM Synthesis & Analysis", level=2)
        doc.add_paragraph(synthesis_text)
        
        doc.add_heading("Audit Metadata", level=2)
        doc.add_paragraph("Validated through Occo AI Semantic Firewall (Verified Context Protocol).")
        
        file_stream = io.BytesIO()
        doc.save(file_stream)
        file_stream.seek(0)
        return file_stream

    @staticmethod
    def generate_pdf(synthesis_text: str, schema_title: str = "Occo AI Compliance Audit") -> io.BytesIO:
        pdf = FPDF()
        pdf.add_page()
        pdf.set_font("Helvetica", "B", 16)
        pdf.cell(0, 10, schema_title, ln=True)
        
        pdf.set_font("Helvetica", "I", 10)
        pdf.cell(0, 8, f"Generated Timestamp: {datetime.utcnow().isoformat()} UTC", ln=True)
        pdf.ln(5)
        
        pdf.set_font("Helvetica", "B", 12)
        pdf.cell(0, 8, "Synthesis & Analysis:", ln=True)
        
        pdf.set_font("Helvetica", "", 10)
        pdf.multi_cell(0, 6, synthesis_text)
        
        pdf_output = pdf.output(dest='S')
        file_stream = io.BytesIO(pdf_output)
        file_stream.seek(0)
        return file_stream

export_service = ExportService()